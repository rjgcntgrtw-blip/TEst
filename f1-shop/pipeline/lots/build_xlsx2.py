"""Собирает f1-shop/data/tmff-assortment.xlsx из lot-variants.csv (варианты лотов из Chrome) + конструкторы."""
from pathlib import Path as _Path

PIPE = str(_Path(__file__).resolve().parents[1]) + "/"  # f1-shop/pipeline/
REPO = str(_Path(__file__).resolve().parents[3]) + "/"  # корень репозитория
import csv
import json
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = REPO + "f1-shop/data/lot-variants.csv"
OUT = REPO + "f1-shop/data/tmff-assortment.xlsx"

F = "Arial"
BLUE = Font(name=F, color="0000FF", size=10)
BLACK = Font(name=F, size=10)
SMALL = Font(name=F, size=9)
BOLD = Font(name=F, size=10, bold=True)
TITLE = Font(name=F, size=14, bold=True)
NOTE = Font(name=F, size=9, italic=True, color="555555")
LINK = Font(name=F, size=10, color="0563C1", underline="single")
YELLOW = PatternFill("solid", fgColor="FFFF00")
SECTION = PatternFill("solid", fgColor="EDEDED")
HEAD = PatternFill("solid", fgColor="1F1F24")
HEADF = Font(name=F, size=10, bold=True, color="FFFFFF")
BORDER = Border(bottom=Side(style="thin", color="D0D0D0"))
TOP = Alignment(vertical="top")
WRAP = Alignment(vertical="top", wrap_text=True)

RU = {"Verstappen": "Ферстаппен", "Tsunoda": "Цунода", "Leclerc": "Леклер", "Hamilton": "Хэмилтон",
      "Piastri": "Пиастри", "Norris": "Норрис", "Russell": "Рассел", "Antonelli": "Антонелли",
      "Alonso": "Алонсо", "Stroll": "Стролл", "Gasly": "Гасли", "Colapinto": "Колапинто",
      "Perez": "Перес", "Sainz": "Сайнс"}
TEAM_ORDER = ["Red Bull", "Ferrari", "Mercedes", "McLaren", "Aston Martin", "Alpine"]
U = lambda i: f"https://aliexpress.ru/item/{i}.html"
C = lambda i: f"https://www.aliexpress.us/item/{i}.html"


def livery_of(row):
    """Ливрея из заметки. У RB21 «Бахрейн» — обычная ливрея, «Япония» — белая (решение владельца)."""
    n = row["заметки"]
    if "вариант 161" in n:
        return "с надписью Leclerc"
    if "ливрея ГП" in n:
        name = n.split("ливрея ГП ")[1].split(" (")[0].strip()
        if name == "Бахрейна":
            return "стандартная"
        if name == "Японии":
            return "белая, ГП Японии"
        return f"ГП {name}"
    return "стандартная"


def removed_reason(row):
    """Варианты, которые владелец решил не брать."""
    n = row["заметки"]
    if "Австралии" in n:
        return "убрано: версия «Австралия»"
    if "Великобритании" in n:
        return "убрано: ливрея «Великобритания»"
    if "Майами 2024" in n and row["машина"].startswith("McLaren"):
        return "убрано: версия McLaren «Майами»"
    return None


LIVERY_NOTE = {
    ("Red Bull RB21", "белая, ГП Японии"): "Белая ливрея в честь Honda (Сузука 2025). Обязательно в ассортименте.",
    ("Ferrari SF-25", "ГП Майами"): "Спецливрея Ferrari на Гран-при Майами.",
    ("Alpine A525", "ГП Монако"): "Другой версии Alpine в лотах нет.",
    ("Red Bull RB19", "ГП США"): "Спецливрея RB19, придуманная фанатами (Остин 2023).",
    ("Ferrari SF-24", "с надписью Leclerc"): "Вариант «SF24-161» — версия с надписью Leclerc. Обычный SF-24 Леклера — отдельной строкой.",
}

# ---------- варианты из CSV ----------
raw = list(csv.DictReader(open(SRC, encoding="utf-8")))
models = [r for r in raw if r["раздел"].startswith("Модели")]
posters = [r for r in raw if r["раздел"] == "Постеры"]

groups = defaultdict(list)
for r in models:
    r["livery"] = livery_of(r)
    r["price"] = int(r["цена_руб"])
    r["sold"] = int(r["заказов"])
    r["key"] = (r["машина"], r["пилот_номер"], r["livery"], r["масштаб"])
    r["removed"] = removed_reason(r)
    if not r["removed"]:
        groups[r["key"]].append(r)

picks = []
for key, rs in groups.items():
    rs.sort(key=lambda r: (r["price"], -r["sold"]))
    best, rest = rs[0], rs[1:]
    best["backup"] = rest
    for r in rest:
        r["dup_of"] = best
    picks.append(best)
print(f"{len(models)} вариантов моделей, убрано {sum(1 for r in models if r['removed'])} -> {len(picks)} уникальных")


def section_of(r):
    if r["масштаб"] == "1:24":
        return "Модели 1:24"
    return "Модели 1:43 — новинки 2026" if r["год"] == "2026" else "Модели 1:43 — сезон 2025"


def team_idx(car):
    return next(i for i, t in enumerate(TEAM_ORDER) if car.startswith(t))


def driver_ru(p):
    name, num = p.split(" #")
    return f"{RU[name]} #{num}"


def site_name(r):
    name = f"{r['машина']} — {driver_ru(r['пилот_номер'])}"
    if r["livery"] != "стандартная":
        name += f", {r['livery']}"
    if r["год"] == "2026":
        name += " · масштаб 1:43"
    return name


SECTIONS = ["Модели 1:43 — сезон 2025", "Модели 1:43 — новинки 2026", "Модели 1:24"]
picks.sort(key=lambda r: (SECTIONS.index(section_of(r)), team_idx(r["машина"]), r["машина"],
                          r["пилот_номер"], r["livery"] != "стандартная", r["livery"]))

# строки главного листа: dict с полями
rows = []
for r in picks:
    notes = []
    if (r["машина"], r["livery"]) in LIVERY_NOTE:
        notes.append(LIVERY_NOTE[(r["машина"], r["livery"])])
    if r["год"] == "2026":
        notes.append("Новинка 2026, оценок ещё нет. На сайте масштаб 1:43 — в названии: модель дороже сезона 2025.")
    backup = "; ".join(f"{b['lot_id'][-6:]} · {b['price']} ₽" for b in r["backup"])
    rows.append(dict(
        section=section_of(r), name=site_name(r), scale=r["масштаб"], price=r["price"],
        car=r["машина"], driver=driver_ru(r["пилот_номер"]), livery=r["livery"], year=int(r["год"]),
        sold=r["sold"], rating=float(r["рейтинг"]) if r["рейтинг"][0].isdigit() else None,
        url=U(r["lot_id"]), variant=r["вариант"], backup=backup, note=" ".join(notes),
        lot_id=r["lot_id"], driver_en=r["пилот_номер"], backups=[(b["lot_id"], b["price"]) for b in r["backup"]],
    ))

# Норрис 1:24 в обычной ливрее: в лотах из Chrome его нет, лот найден поиском aliexpress.ru (проверить ливрею)
NORRIS = dict(
    section="Модели 1:24", name="McLaren MCL38 — Норрис #4", scale="1:24", price=1842, car="McLaren MCL38",
    driver="Норрис #4", livery="стандартная", year=2024, sold=150, rating=5.0, url=U("1005008619594051"),
    variant="Норрис #4 — выбрать в лоте", backup="009441469506 · 2 699 ₽",
    note=("Лот найден поиском aliexpress.ru, в CSV из Chrome его нет: проверить по фото, что ливрея обычная, а не «Майами», "
          "и цену варианта. Запасной: 1005009441469506 · 2 699 ₽, 503 заказа, 4.9."),
    lot_id="1005008619594051", driver_en="Norris #4", backups=[("1005009441469506", 2699)], check_livery=True,
)
i = next(k for k, d in enumerate(rows) if d["car"] == "McLaren MCL38")
rows.insert(i, NORRIS)

# конструкторы: данные поиска aliexpress.com (цены пересчитаны из $), вариант с коробкой не проверен
KITS = [
    ("Конструкторы — лицензия CaDA", "CaDA BWT Alpine A525", "Alpine A525", "1:24", 4524, 268, 4.9, C("3256811986486740"), 2025,
     "", "Официальная лицензия Alpine. Цена — верхняя оценка (зачёркнутая цена), реальная может быть до двух раз ниже."),
    ("Конструкторы — лицензия CaDA", "CaDA Kick Sauber C44", "Kick Sauber C44", "1:24", 6648, 340, 4.7, C("3256807694767657"), 2024,
     "", "Официальная лицензия Sauber. Цена — верхняя оценка (зачёркнутая цена), реальная может быть до двух раз ниже."),
]
for sec, name, car, scale, price, sold, rating, url, year, backup, note in KITS:
    rows.append(dict(section=sec, name=name, scale=scale, price=price, car=car, driver="—",
                     livery="реплика" if "реплика" in sec else "CaDA", year=year, sold=sold, rating=rating,
                     url=url, variant=None, backup=backup, note=note, check_box=True))
# реплики 1:8 — подборка от 25.09.2026 (site/make_kits.py)
for k in json.load(open(PIPE + "lots/kits_pick.json", encoding="utf-8")):
    rows.append(dict(section=k["section"], name=k["name"], scale="1:8", price=k["price"], car=k["car"], driver="—",
                     livery="реплика", year=k["year"], sold=k["sold"], rating=k["rating"], url=k["url"],
                     variant=k["variant"], backup=k["backup"], note=k["note"], check_box=True))

# постеры: по строке на дизайн, размер по умолчанию 30×40
POSTER_YEAR = {"McLaren M23": 1974, "McLaren-Mercedes MP4-23": 2008, "McLaren MCL34": 2019, "Lotus 98T": 1986,
               "Red Bull Racing RB16": 2020, "Lotus 99T": 1987, "Jordan 191": 1991, "Mercedes AMG W10": 2019,
               "McLaren MP4/5B": 1990, "McLaren MP4/4": 1988, "McLaren MP4/6": 1991, "Brawn BGP001": 2009,
               "Ferrari 643": 1991, "Renault R25": 2005, "Ferrari F2007": 2007, "Toro Rosso STR3": 2008,
               "Toro Rosso STR14": 2019, "Ferrari 312T": 1975, "Toleman TG184": 1984, "Red Bull Racing RB7": 2011,
               "Schumacher / Ferrari (японский стиль)": None}
POSTER_SIZE = "30x40 см"
seen = []
for r in posters:
    if r["масштаб"] != POSTER_SIZE or r["машина"] in seen:
        continue
    seen.append(r["машина"])
    design = r["машина"].replace("Schumacher / Ferrari (японский стиль)", "Шумахер, Ferrari — японский стиль")
    note = "Холст без рамки. Цены других размеров — лист «Постеры — размеры»."
    rows.append(dict(section="Постеры", name=f"Постер {design}", scale="30×40 см", price=int(r["цена_руб"]),
                     car=design, driver="—", livery="—", year=POSTER_YEAR[r["машина"]], sold=int(r["заказов"]),
                     rating=float(r["рейтинг"]), url=U(r["lot_id"]), variant=r["вариант"], backup="", note=note))
print("постеров:", len(seen))

# ---------- книга ----------
wb = Workbook()

# Параметры
p = wb.active
p.title = "Параметры"
params = [
    ("Порог 1: закупка дешевле, ₽", 1500, "Правило наценки из f1-shop/01-business-model.md"),
    ("Наценка до порога 1, ×", 2.0, "Дешёвые позиции: ×2"),
    ("Порог 2: закупка до, ₽", 7000, ""),
    ("Наценка между порогами, ×", 1.6, "Средние позиции: ×1,6"),
    ("Наценка дороже порога 2, ×", 1.35, "Дорогие позиции: ×1,35"),
    ("Эквайринг, доля", 0.035, "ЮKassa, карта; по СБП меньше (f1-shop/03-payments.md)"),
    ("Налог УСН «Доходы», доля", 0.06, "ИП на УСН 6 %"),
    ("Резерв на возвраты, доля", 0.05, "f1-shop/01-business-model.md"),
    ("Лимит цены для 1:24, ₽", 2800, "Правило владельца магазина: 1:24 берём до 2 800 ₽"),
    ("Курс доллара, ₽", 84.2, "ЦБ РФ на 20.09.2026 — для лотов с aliexpress.com (конструкторы, мерч)"),
    ("Резерв на возвраты, одежда, доля", 0.10, "По одежде возвратов больше — размер (f1-shop/10-content.md)"),
]
p["A1"] = "Параметры расчёта"; p["A1"].font = TITLE
p["A2"] = "Синие ячейки можно менять — цены и прибыль на листах «Ассортимент» и «Постеры — размеры» пересчитаются."
p["A2"].font = NOTE
for i, (name, val, note) in enumerate(params, start=4):
    p.cell(i, 1, name).font = BLACK
    c = p.cell(i, 2, val); c.font = BLUE
    c.number_format = "0.0%" if isinstance(val, float) and val < 1 else ("0.00" if isinstance(val, float) else "#,##0")
    p.cell(i, 3, note).font = NOTE
p.column_dimensions["A"].width = 36; p.column_dimensions["B"].width = 12; p.column_dimensions["C"].width = 80
T1, M1, T2, M2, M3, ACQ, TAX, RES, LIM24 = [f"Параметры!$B${i}" for i in range(4, 13)]
RES_WEAR = "Параметры!$B$14"


def our_price(ref):
    return (f'IF({ref}="","",IF({ref}<{T1},CEILING({ref}*{M1},100)-10,'
            f'IF({ref}<={T2},CEILING({ref}*{M2},100)-10,CEILING({ref}*{M3},100)-10)))')


def profit(price_ref, ours_ref):
    return f'IF({price_ref}="","",ROUND({ours_ref}*(1-{ACQ}-{TAX}-{RES})-{price_ref},0))'


# Ассортимент
s = wb.create_sheet("Ассортимент", 0)
cols = [("Раздел", 24), ("Название на сайте", 44), ("Масштаб", 9), ("Цена на Ali, ₽", 11), ("Наша цена, ₽", 11),
        ("Прибыль с шт., ₽", 11), ("Болид", 20), ("Пилот", 15), ("Ливрея / версия", 17), ("Год", 6),
        ("Заказов лота", 9), ("Рейтинг", 8), ("Лот", 9), ("Вариант на Ali", 17), ("Запасной лот · цена", 17),
        ("Повтор?", 9), ("Лимит 1:24", 10), ("Заметки и что проверить", 70)]
NC = len(cols)
s["A1"] = "TMFF — стартовый ассортимент"; s["A1"].font = TITLE
s["A2"] = ("Синим — данные AliExpress (варианты лотов собраны в Chrome 25.09.2026, конструкторы — из поиска aliexpress.com), "
           "чёрным — формулы, жёлтым — проверить. Цена — за товар, без платной доставки, если она есть. "
           "У конструкторов с aliexpress.com жёлтая цена — верхняя оценка: реальная может быть до двух раз ниже.")
s["A2"].font = NOTE
s["A3"] = ("Повторы убраны: болид + пилот + ливрея + масштаб встречаются один раз, по самой низкой цене. "
           "Другие лоты с тем же вариантом — в «Запасной лот». «Вариант на Ali» — что выбрать при заказе.")
s["A3"].font = NOTE
HR = 4
for j, (h, w) in enumerate(cols, start=1):
    c = s.cell(HR, j, h); c.font = HEADF; c.fill = HEAD; c.alignment = Alignment(vertical="center", wrap_text=True)
    s.column_dimensions[get_column_letter(j)].width = w
s.row_dimensions[HR].height = 30

first = HR + 1
last = first + len(rows) - 1
for k, d in enumerate(rows):
    r = first + k
    vals = [d["section"], d["name"], d["scale"], d["price"], None, None, d["car"], d["driver"], d["livery"], d["year"],
            d["sold"], d["rating"], None, d["variant"], d["backup"], None, None, d["note"]]
    for j, v in enumerate(vals, start=1):
        c = s.cell(r, j, v)
        c.font = BLACK
        c.alignment = TOP
        c.border = BORDER
    s.cell(r, 2).font = BOLD
    s.cell(r, 2).alignment = WRAP
    s.cell(r, 1).alignment = WRAP
    for j in (3, 4, 11, 12, 14):
        s.cell(r, j).font = BLUE
    s.cell(r, 4).number_format = "#,##0"
    s.cell(r, 5).value = "=" + our_price(f"D{r}")
    s.cell(r, 5).number_format = "#,##0"
    s.cell(r, 6).value = "=" + profit(f"D{r}", f"E{r}")
    s.cell(r, 6).number_format = "#,##0"
    s.cell(r, 11).number_format = '#,##0"+"' if d["sold"] == 1000 else "#,##0"
    s.cell(r, 12).number_format = "0.0"
    c = s.cell(r, 13, "открыть"); c.hyperlink = d["url"]; c.font = LINK
    s.cell(r, 15).font = SMALL
    s.cell(r, 16).value = (f'=IF(COUNTIFS($G${first}:$G${last},G{r},$H${first}:$H${last},H{r},'
                           f'$I${first}:$I${last},I{r},$C${first}:$C${last},C{r})>1,"повтор","")')
    s.cell(r, 17).value = f'=IF(C{r}<>"1:24","",IF(D{r}="","проверить",IF(D{r}<={LIM24},"ок","дороже лимита")))'
    s.cell(r, 18).font = SMALL
    s.cell(r, 18).alignment = WRAP
    if d.get("check_box"):
        s.cell(r, 14).value = "проверить"
        s.cell(r, 14).fill = YELLOW
        s.cell(r, 18).value = (d["note"] + " Вариант «с коробкой» и его цену не проверяли: в CSV из Chrome этого лота нет.").strip()
        if "верхняя оценка" in d["note"]:
            s.cell(r, 4).fill = YELLOW
    if d["rating"] is None and not d["section"].startswith("Модели 1:43 — новинки"):
        s.cell(r, 12).fill = YELLOW
    needs_photo = "сверить по фото" in d["note"] or d.get("check_livery")
    if needs_photo:
        s.cell(r, 9).fill = YELLOW

# цены конструкторов пересчитаны из $
for r in range(first, last + 1):
    if "aliexpress.us" in s.cell(r, 13).hyperlink.target:
        s.cell(r, 4).comment = Comment("Обычная цена из выдачи aliexpress.com (без скидки для новых покупателей), "
                                       "пересчитана из $ по курсу 84,2 ₽/$.", "TMFF")

# подсветка повторов и превышения лимита
red = PatternFill("solid", fgColor="F4CCCC")
s.conditional_formatting.add(f"P{first}:Q{last}", FormulaRule(formula=[f'OR(P{first}="повтор",Q{first}="дороже лимита")'], fill=red))

# итоги
tot = last + 2
s.cell(tot, 1, "Итого").font = BOLD
summary = [("Позиций всего", f"=COUNTA(B{first}:B{last})", "#,##0")]
for sec in SECTIONS + ["Конструкторы — лицензия CaDA", "Конструкторы — реплика", "Постеры"]:
    summary.append((sec, f'=COUNTIF($A${first}:$A${last},"{sec}")', "#,##0"))
summary += [
    ("Повторов", f'=COUNTIF(P{first}:P{last},"повтор")', "#,##0"),
    ("1:24 дороже лимита", f'=COUNTIF(Q{first}:Q{last},"дороже лимита")', "#,##0"),
    ("Средняя прибыль с шт., ₽", f"=IFERROR(ROUND(AVERAGE(F{first}:F{last}),0),\"\")", "#,##0"),
    ("Средняя прибыль, модели 1:43, ₽", f'=IFERROR(ROUND(AVERAGEIFS(F{first}:F{last},C{first}:C{last},"1:43"),0),"")', "#,##0"),
    ("Средняя прибыль, модели 1:24, ₽",
     f'=IFERROR(ROUND(AVERAGEIFS(F{first}:F{last},C{first}:C{last},"1:24",A{first}:A{last},"Модели 1:24"),0),"")', "#,##0"),
]
for i, (label, formula, fmt) in enumerate(summary):
    s.cell(tot + 1 + i, 1, label).font = BLACK
    c = s.cell(tot + 1 + i, 2, formula); c.font = BOLD; c.number_format = fmt; c.alignment = Alignment(horizontal="left")
s.freeze_panes = s.cell(first, 3)
s.auto_filter.ref = f"A{HR}:{get_column_letter(NC)}{last}"

# Мерч и кепки
MERCH = PIPE + "merch/"
mp = json.load(open(MERCH + "picks.json", encoding="utf-8"))
TEAM_RU = {"": "—", "ferrari": "Ferrari", "redbull": "Red Bull", "mclaren": "McLaren", "mercedes": "Mercedes",
           "aston": "Aston Martin", "williams": "Williams", "audi": "Audi", "haas": "Haas", "alpine": "Alpine", "legend": "Легенды"}
m = wb.create_sheet("Мерч и кепки", 1)
m["A1"] = "TMFF — одежда в стиле командной формы"; m["A1"].font = TITLE
m["A2"] = ("Поиск aliexpress.com 25.09.2026. Цена пересчитана из $ по курсу с листа «Параметры»; жёлтая — верхняя оценка "
           "(в выдаче была только цена для новичков), реальная может быть до двух раз ниже. Прибыль — с резервом на возвраты по одежде (10 %).")
m["A2"].font = NOTE
m["A3"] = ("Только вещи в стиле официальной формы команд, на сайте — с пометкой «Реплика». Командные кепки и худи Ferrari, Red Bull, "
           "Mercedes и McLaren на aliexpress.com почти не показываются — стоит поискать их на aliexpress.ru. Номер № — на картинке merch-candidates.jpg.")
m["A3"].font = NOTE
mcols = [("№", 5), ("Команда", 16), ("Название на сайте", 38), ("Вещь", 10), ("Команда", 12), ("Цена на Ali, ₽", 11),
         ("Наша цена, ₽", 11), ("Прибыль с шт., ₽", 11), ("Заказов", 9), ("Рейтинг", 8), ("Лот", 9), ("Заметки и что проверить", 80)]
for j, (h, w) in enumerate(mcols, start=1):
    c = m.cell(HR, j, h); c.font = HEADF; c.fill = HEAD; c.alignment = Alignment(vertical="center", wrap_text=True)
    m.column_dimensions[get_column_letter(j)].width = w
m.row_dimensions[HR].height = 30
mf = HR + 1
for k, d in enumerate(mp):
    r = mf + k
    vals = [k + 1, d["section"], d["name"], d["kind"], TEAM_RU.get(d["team"], d["team"]), d["price"], None, None,
            d["sold"], d["rating"], None, d["note"]]
    for j, v in enumerate(vals, start=1):
        c = m.cell(r, j, v); c.font = BLACK; c.alignment = TOP; c.border = BORDER
    m.cell(r, 2).alignment = WRAP
    m.cell(r, 3).font = BOLD; m.cell(r, 3).alignment = WRAP
    for j in (6, 9, 10):
        m.cell(r, j).font = BLUE
    m.cell(r, 6).number_format = "#,##0"
    m.cell(r, 6).comment = Comment(f"${d['usd']} × курс" + ("" if d.get("exact", True) else "; верхняя оценка — зачёркнутая цена"), "TMFF")
    m.cell(r, 7).value = "=" + our_price(f"F{r}"); m.cell(r, 7).number_format = "#,##0"
    m.cell(r, 8).value = f'=IF(F{r}="","",ROUND(G{r}*(1-{ACQ}-{TAX}-{RES_WEAR})-F{r},0))'; m.cell(r, 8).number_format = "#,##0"
    m.cell(r, 9).number_format = "#,##0"; m.cell(r, 10).number_format = "0.0"
    c = m.cell(r, 11, "открыть"); c.hyperlink = d["url"]; c.font = LINK
    m.cell(r, 12).font = SMALL; m.cell(r, 12).alignment = WRAP
    if not d.get("exact", True):
        m.cell(r, 6).fill = YELLOW
    if d["rating"] is not None and d["rating"] < 4.5:
        m.cell(r, 10).fill = YELLOW
    if d["sold"] < 50:
        m.cell(r, 9).fill = YELLOW
ml = mf + len(mp) - 1
mt = ml + 2
m.cell(mt, 1, "Итого").font = BOLD
msum = [("Позиций", f"=COUNTA(C{mf}:C{ml})")]
for kind_ru in ("футболка", "кофта", "кепка"):
    msum.append((kind_ru, f'=COUNTIF($D${mf}:$D${ml},"{kind_ru}")'))
msum.append(("Средняя прибыль, ₽", f'=IFERROR(ROUND(AVERAGE($H${mf}:$H${ml}),0),"")'))
for i, (label, formula) in enumerate(msum):
    m.cell(mt + 1 + i, 2, label).font = BLACK
    c = m.cell(mt + 1 + i, 3, formula); c.font = BOLD; c.number_format = "#,##0"; c.alignment = Alignment(horizontal="left")
m.freeze_panes = m.cell(mf, 4)
m.auto_filter.ref = f"A{HR}:{get_column_letter(len(mcols))}{ml}"

# Мерч — все лоты
allm = [r for r in json.load(open(MERCH + "items.json", encoding="utf-8")) if r["kind"] in ("cap", "tshirt", "hoodie", "polo")]
KIND_RU = {"cap": "кепка", "tshirt": "футболка", "hoodie": "кофта", "polo": "поло"}
a = wb.create_sheet("Мерч — все лоты")
a["A1"] = f"Все найденные лоты одежды и кепок — {len(allm)}"; a["A1"].font = TITLE
a["A2"] = ("Тип, команда и риск определены по названию лота автоматически — могут быть ошибки. Отсортировано по заказам. "
           "Фильтр в заголовке — чтобы смотреть по вещи, команде или риску.")
a["A2"].font = NOTE
ah = [("Вещь", 9), ("Команда", 12), ("Риск по названию", 24), ("Название лота", 90), ("Цена, ₽", 9), ("Для новичков, $", 10),
      ("Заказов", 9), ("Рейтинг", 8), ("Лот", 9), ("Запрос", 28)]
for j, (h, w) in enumerate(ah, start=1):
    c = a.cell(4, j, h); c.font = HEADF; c.fill = HEAD
    a.column_dimensions[get_column_letter(j)].width = w
for i, r in enumerate(allm, start=5):
    vals = [KIND_RU[r["kind"]], TEAM_RU.get(r["team"], r["team"] or "—"), r["risk"], r["title"], r["price"], r["promo_usd"],
            r["sold"], r["rating"], "открыть", r["query"]]
    for j, v in enumerate(vals, start=1):
        c = a.cell(i, j, v); c.font = BLACK
    a.cell(i, 5).number_format = "#,##0"; a.cell(i, 7).number_format = "#,##0"
    a.cell(i, 9).hyperlink = f"https://www.aliexpress.us/item/{r['id']}.html"; a.cell(i, 9).font = LINK
a.freeze_panes = "A5"
a.auto_filter.ref = f"A4:J{4 + len(allm)}"

# Постеры — размеры
ps = wb.create_sheet("Постеры — размеры")
ps["A1"] = "Постеры: цена по размерам"; ps["A1"].font = TITLE
ps["A2"] = ("Лот " + posters[0]["lot_id"] + ": 21 дизайн, холст без рамки, цена одинакова для всех дизайнов. "
            "Оценка 4.4, 112 заказов. «На сайт» — какие размеры предлагать (по умолчанию 4 ходовых).")
ps["A2"].font = NOTE
hdr = ["Размер", "Цена на Ali, ₽", "Наша цена, ₽", "Прибыль с шт., ₽", "На сайт"]
for j, h in enumerate(hdr, start=1):
    c = ps.cell(4, j, h); c.font = HEADF; c.fill = HEAD
sizes = []
for r in posters:
    if r["машина"] == posters[0]["машина"]:
        sizes.append((r["масштаб"].replace("x", "×"), int(r["цена_руб"])))
ON_SITE = {"20×30 см", "30×40 см", "40×60 см", "50×70 см"}
for i, (size, price) in enumerate(sizes, start=5):
    ps.cell(i, 1, size).font = BLACK
    c = ps.cell(i, 2, price); c.font = BLUE; c.number_format = "#,##0"
    c = ps.cell(i, 3, "=" + our_price(f"B{i}")); c.font = BLACK; c.number_format = "#,##0"
    c = ps.cell(i, 4, "=" + profit(f"B{i}", f"C{i}")); c.font = BLACK; c.number_format = "#,##0"
    c = ps.cell(i, 5, "да" if size in ON_SITE else ""); c.font = BLUE
for col, w in zip("ABCDE", (12, 14, 14, 16, 10)):
    ps.column_dimensions[col].width = w
pn = 5 + len(sizes) + 1
ps.cell(pn, 1, "Дизайны").font = BOLD
for i, name in enumerate(seen, start=pn + 1):
    sku = next(r["вариант"].split(" / ")[0] for r in posters if r["машина"] == name)
    ps.cell(i, 1, sku).font = BLUE
    ps.cell(i, 2, name).font = BLACK

# Варианты AliExpress (все 57 вариантов моделей)
v = wb.create_sheet("Варианты AliExpress")
v["A1"] = "Все варианты моделей из лотов (CSV из Chrome, 25.09.2026)"; v["A1"].font = TITLE
v["A2"] = ("«В ассортименте» — вариант взят на лист «Ассортимент». «Дубль» — такой же есть дешевле или с бо́льшим числом заказов. "
           "«Убрано» — версии, которые решили не брать.")
v["A2"].font = NOTE
vh = [("Лот", 18), ("Вариант на Ali", 20), ("Болид", 20), ("Пилот", 15), ("Ливрея / версия", 17), ("Год", 6),
      ("Масштаб", 9), ("Цена, ₽", 10), ("Рейтинг", 9), ("Заказов", 9), ("Статус", 34)]
for j, (h, w) in enumerate(vh, start=1):
    c = v.cell(4, j, h); c.font = HEADF; c.fill = HEAD
    v.column_dimensions[get_column_letter(j)].width = w
for i, r in enumerate(models, start=5):
    if r["removed"]:
        status = r["removed"]
    elif "dup_of" in r:
        status = f"дубль: берём лот …{r['dup_of']['lot_id'][-6:]} за {r['dup_of']['price']} ₽"
    else:
        status = "в ассортименте"
    vals = [r["lot_id"], r["вариант"], r["машина"], driver_ru(r["пилот_номер"]), r["livery"], int(r["год"]), r["масштаб"],
            r["price"], r["рейтинг"], r["sold"], status]
    for j, val in enumerate(vals, start=1):
        c = v.cell(i, j, val); c.font = BLACK; c.border = BORDER
    v.cell(i, 1).hyperlink = U(r["lot_id"]); v.cell(i, 1).font = LINK
    v.cell(i, 8).number_format = "#,##0"
    if status == "в ассортименте":
        v.cell(i, 11).font = BOLD
v.freeze_panes = "A5"
v.auto_filter.ref = f"A4:K{4 + len(models)}"

# Что уточнить
q = wb.create_sheet("Что уточнить")
q["A1"] = "Что осталось проверить"; q["A1"].font = TITLE
items = [
    ("Конструкторы CaDA 1:24 (2 лота)", "Цена нужного варианта, официальный магазин CaDA, число деталей и возраст на коробке. "
     "В CSV из Chrome этих лотов нет, цены — из поиска aliexpress.com: там была только цена для новых покупателей, "
     "в таблице — зачёркнутая как верхняя оценка. Реальная может быть до двух раз ниже."),
    ("Реплики 1:8 (7 машин, 5 лотов)", "Цена нужного варианта и вариант «с коробкой»; фото коробки в отзывах: без логотипа LEGO, печать нормального качества. "
     "Жёлтые цены — верхняя оценка."),
    ("Доставка", "Цены в таблице — за товар. Если доставка в РФ платная, прибавить её к цене на Ali."),
    ("Новинки 2026 (1:43)", "Оценок нет, 31–45 заказов. Сначала заказать по одной для проверки качества."),
    ("Мерч в стиле командной формы", "Цена нужного размера (жёлтые — верхняя оценка), фото из отзывов: качество печати и вышивки. "
     "Кепки и худи Ferrari, Red Bull, Mercedes, McLaren поискать на aliexpress.ru — на aliexpress.com их почти нет."),
    ("Мерч — размеры", "У каждого продавца одежды взять таблицу размеров в сантиметрах: азиатская сетка обычно на 1–2 размера меньше."),
    ("Постеры", "Оценка лота 4.4 — ниже остальных; посмотреть фото печати в отзывах."),
]
for i, (a, b) in enumerate(items, start=3):
    q.cell(i, 1, a).font = BOLD
    q.cell(i, 1).alignment = WRAP
    c = q.cell(i, 2, b); c.font = BLACK; c.alignment = WRAP
q.column_dimensions["A"].width = 30; q.column_dimensions["B"].width = 110

json.dump(rows, open(PIPE + "lots/assortment.json", "w",
                     encoding="utf-8"), ensure_ascii=False, indent=1)
wb.move_sheet("Параметры", offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("Параметры"))
wb.save(OUT)
print("saved", OUT, "rows", len(rows), "first", first, "last", last)
